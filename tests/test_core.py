"""Tests — run with: pytest  (no API keys needed, uses mock slots)."""
import asyncio

import pytest

from theosis.config import demo_aggregator, demo_slots
from theosis.core import audit_pairs, theosis
from theosis.metrics import avg_similarity, estimate_cost, verdict_score
from theosis.models import ModelSlot


# ── pipeline basics ─────────────────────────────────────────────────────────
def test_audit_pairs_is_round_robin():
    class S:
        def __init__(self, name):
            self.name = name

    a, b, c = S("a"), S("b"), S("c")
    pairs = [(t.name, au.name) for t, au in audit_pairs([a, b, c])]
    assert pairs == [("a", "b"), ("b", "c"), ("c", "a")]


def test_pipeline_runs_in_mock_mode():
    slots = demo_slots()
    for s in slots:
        s.enabled = True
    final, trail = asyncio.run(theosis("xin chào", slots, demo_aggregator(), max_rounds=1))
    assert isinstance(final, str) and final.strip()
    assert trail["fanout"] and len(trail["rounds"]) == 1
    for s in slots:
        assert s.name in trail["fanout"]


def test_single_slot_skips_cross_audit():
    slots = demo_slots()
    for s in slots:
        s.enabled = s.name == "opus"
    final, trail = asyncio.run(theosis("hi", slots, demo_aggregator(), max_rounds=2))
    assert trail["rounds"] == []
    assert final.strip()


def test_no_enabled_slot_raises():
    slots = demo_slots()
    for s in slots:
        s.enabled = False
    try:
        asyncio.run(theosis("hi", slots, demo_aggregator(), max_rounds=1))
        assert False, "expected ValueError"
    except ValueError:
        pass


# ── V2: resilience ──────────────────────────────────────────────────────────
def test_failing_slot_is_dropped_not_fatal():
    good = ModelSlot("good", "m", "mock://")
    bad = ModelSlot("bad", "m", "http://127.0.0.1:1/v1", api_key="x")  # connection refused
    final, trail = asyncio.run(theosis("hi", [good, bad], demo_aggregator(), max_rounds=1))
    assert "good" in trail["fanout"]
    assert "bad" not in trail["fanout"]  # bị loại, không làm sập mẻ
    assert final.strip()


# ── V2: cost meter ──────────────────────────────────────────────────────────
def test_cost_meter_accumulates():
    slots = demo_slots()
    for s in slots:
        s.enabled = True
    _, trail = asyncio.run(theosis("hi", slots, demo_aggregator(), max_rounds=1))
    cost = trail["cost"]
    assert cost["calls"] > 0 and cost["total_tokens"] > 0


def test_estimate_cost_known_and_default():
    # gpt-4o: 2.5 in / 10 out per 1M
    c = estimate_cost("gpt-4o", {"prompt_tokens": 1_000_000, "completion_tokens": 1_000_000})
    assert abs(c - 12.5) < 1e-6
    # model lạ -> bảng giá mặc định
    c2 = estimate_cost("unknown-model", {"prompt_tokens": 1_000_000, "completion_tokens": 0})
    assert c2 > 0


# ── V2: convergence + scoring helpers ───────────────────────────────────────
def test_avg_similarity_identical_is_one():
    assert avg_similarity({"a": "hello world"}, {"a": "hello world"}) == 1.0


def test_verdict_score_mapping():
    assert verdict_score("VERDICT: strong") == 1.0
    assert verdict_score("VERDICT: weak") == 0.2
    assert verdict_score("no verdict here") == 0.5


def test_scores_present_in_trail():
    slots = demo_slots()
    for s in slots:
        s.enabled = True
    _, trail = asyncio.run(theosis("hi", slots, demo_aggregator(), max_rounds=1))
    assert set(trail["scores"]) >= {"opus", "gpt"}


# ── V2: multi-auditor (M-of-N) ──────────────────────────────────────────────
def test_default_single_auditor_per_answer():
    slots = demo_slots()
    for s in slots:
        s.enabled = True
    _, trail = asyncio.run(theosis("hi", slots, demo_aggregator(), max_rounds=1))
    rnd = trail["rounds"][0]
    assert all(len(revs) == 1 for revs in rnd["reviews"].values())


def test_multi_auditor_gives_k_reviews():
    slots = demo_slots()  # 3 slots
    for s in slots:
        s.enabled = True
    _, trail = asyncio.run(theosis("hi", slots, demo_aggregator(), max_rounds=1, auditors_per_answer=2))
    rnd = trail["rounds"][0]
    assert len(rnd["reviews"]) == 3
    assert all(len(revs) == 2 for revs in rnd["reviews"].values())


def test_audit_assignments_clamps_to_available():
    from theosis.core import audit_assignments
    from theosis.models import ModelSlot

    slots = [ModelSlot(n, "m", "mock://") for n in ("a", "b")]
    asg = audit_assignments(slots, 5)  # xin 5 nhưng chỉ 2 slot -> clamp về 1
    assert all(len(auditors) == 1 for _, auditors in asg)


def test_pricing_override_via_set_cost_table():
    from theosis import metrics

    metrics.set_cost_table({"my-model": [10.0, 20.0]})
    c = metrics.estimate_cost("my-model", {"prompt_tokens": 1_000_000, "completion_tokens": 0})
    assert abs(c - 10.0) < 1e-6


# ── V2.2: executor (ground-truth verifiers) ─────────────────────────────────
def test_verifier_classifies_code():
    from theosis.verifiers import run_verifiers

    assert run_verifiers("```python\nprint(1 + 1)\n```")["status"] == "pass"
    assert run_verifiers("```python\nraise ValueError('x')\n```")["status"] == "fail"
    assert run_verifiers("chỉ là văn xuôi, không code")["status"] == "na"


def test_verifier_checks_arithmetic():
    from theosis.verifiers import run_verifiers

    assert run_verifiers("Khẳng định: 2 + 2 = 5")["status"] == "fail"
    assert run_verifiers("Đúng là 3 * 4 = 12")["status"] == "pass"


def test_safe_eval_rejects_non_arithmetic():
    from theosis.verifiers import _safe_eval

    for bad in ["__import__('os')", "x + 1", "open('f')"]:
        try:
            _safe_eval(bad)
            assert False, f"phải từ chối: {bad}"
        except Exception:
            pass


def test_executor_adds_evidence_to_trail():
    slots = demo_slots()
    for s in slots:
        s.enabled = True
    _, trail = asyncio.run(theosis("hi", slots, demo_aggregator(), max_rounds=1, use_executor=True))
    assert "evidence" in trail and isinstance(trail["evidence"], dict)


# ── V2.3: pluggable strategies ──────────────────────────────────────────────
def test_strategy_all_vs_all():
    from theosis.models import ModelSlot
    from theosis.strategies import all_vs_all

    s = [ModelSlot(n, "m", "mock://") for n in ("a", "b", "c")]
    asg = {t.name: [a.name for a in auds] for t, auds in all_vs_all(s, 1)}
    assert asg["a"] == ["b", "c"] and asg["b"] == ["a", "c"] and asg["c"] == ["a", "b"]


def test_strategy_star():
    from theosis.models import ModelSlot
    from theosis.strategies import star

    s = [ModelSlot(n, "m", "mock://") for n in ("a", "b", "c")]
    asg = {t.name: [a.name for a in auds] for t, auds in star(s, 1)}
    assert asg["b"] == ["a"] and asg["c"] == ["a"]  # giám khảo 'a' soi mọi câu khác
    assert asg["a"] == ["b"]                          # 'a' do á quân 'b' soi


def test_get_strategy_fallback():
    from theosis.strategies import get_strategy, round_robin

    assert get_strategy("khong-ton-tai") is round_robin


def test_all_vs_all_integration():
    slots = demo_slots()  # 3 slots
    for s in slots:
        s.enabled = True
    _, trail = asyncio.run(theosis("hi", slots, demo_aggregator(), max_rounds=1, strategy="all_vs_all"))
    rnd = trail["rounds"][0]
    assert all(len(revs) == 2 for revs in rnd["reviews"].values())


# ── V2.4: streaming ─────────────────────────────────────────────────────────
def test_theosis_stream_yields_final():
    from theosis.core import theosis_stream

    slots = demo_slots()
    for s in slots:
        s.enabled = True

    async def collect():
        return "".join([t async for t in theosis_stream("hi", slots, demo_aggregator(), max_rounds=1)])

    text = asyncio.run(collect())
    assert text.strip() and "mock" in text  # merge mock được stream ra


def test_theosis_stream_single_slot():
    from theosis.core import theosis_stream

    slots = demo_slots()
    for s in slots:
        s.enabled = s.name == "opus"

    async def collect():
        return "".join([t async for t in theosis_stream("hi", slots, demo_aggregator(), max_rounds=1)])

    assert asyncio.run(collect()).strip()


# ── V3 / Phase A: smart router ──────────────────────────────────────────────
def test_router_parse_validates_and_clamps():
    from theosis.router import parse_plan
    from theosis.models import ModelSlot

    sl = [ModelSlot(n, "m", "mock://") for n in ("a", "b", "c")]
    p = parse_plan('{"slots":["a","GHOST","b"],"strategy":"bad","rounds":99,"use_executor":"true"}', sl)
    assert p.slots == ["a", "b"]          # bỏ slot ma
    assert p.strategy == "round_robin"    # strategy bậy -> default
    assert p.rounds == 3                  # clamp
    assert p.use_executor is True
    assert p.routed is True


def test_router_fallback_on_garbage():
    from theosis.router import parse_plan
    from theosis.models import ModelSlot

    sl = [ModelSlot(n, "m", "mock://") for n in ("a", "b")]
    p = parse_plan("không có json", sl, def_strategy="star", def_rounds=2)
    assert p.routed is False
    assert p.slots == ["a", "b"]          # giữ tất cả
    assert p.strategy == "star"           # giữ default


def test_router_selects_subset_and_records_trail():
    slots = demo_slots()  # 3 slot
    for s in slots:
        s.enabled = True
    _, trail = asyncio.run(theosis("giải thích đệ quy", slots, demo_aggregator(), use_router=True))
    r = trail["route"]
    assert r is not None and r["routed"]
    assert len(r["slots"]) == 2                       # router (mock) chọn 2
    assert set(trail["fanout"]) == set(r["slots"])    # chỉ chạy slot được chọn


def test_router_off_is_backward_compatible():
    slots = demo_slots()
    for s in slots:
        s.enabled = True
    _, trail = asyncio.run(theosis("hi", slots, demo_aggregator(), max_rounds=1))
    assert trail["route"] is None
    assert set(trail["fanout"]) == {"opus", "gpt", "deepseek"}


# ── V3 / Phase B: immune memory ─────────────────────────────────────────────
def test_memory_store_crud(tmp_path):
    from theosis.memory import MemoryStore, Rule

    p = tmp_path / "mem.json"
    s = MemoryStore(str(p))
    r = s.add_rule(Rule(guidance="Xử lý base case", task_type="code", keywords=["đệ quy"]))
    assert len(s.all_rules()) == 1 and r.id.startswith("r_")
    assert MemoryStore(str(p)).all_rules()[0].guidance == "Xử lý base case"   # reload từ đĩa
    assert s.relevant(task_type="code") and s.relevant(task_type="creative") == []
    assert s.remove(r.id) and s.all_rules() == []


def test_parse_rule_variants():
    from theosis.memory import parse_rule

    r = parse_rule('{"guidance":"G","task_type":"math","keywords":["a"]}')
    assert r.guidance == "G" and r.task_type == "math" and r.keywords == ["a"]
    assert parse_rule("chỉ là text").guidance == "chỉ là text"
    assert parse_rule("   ") is None


def test_log_metrics_is_content_free(tmp_path):
    from theosis.memory import log_metrics

    secret = "ANSWER_SECRET_42"
    trail = {"route": {"task_type": "code", "strategy": "round_robin"},
             "fanout": {"opus": secret}, "rounds": [1], "stopped_reason": "converged",
             "cost": {"total_tokens": 10}, "scores": {"opus": 0.9}}
    p = tmp_path / "m.jsonl"
    rec = log_metrics(trail, str(p))
    assert rec["task_type"] == "code" and rec["n_models"] == 1
    assert secret not in p.read_text(encoding="utf-8")        # content-free trên đĩa


def test_make_rule_and_injection_are_content_free(tmp_path):
    import httpx

    from theosis.core import make_rule, theosis
    from theosis.memory import MemoryStore

    slots = demo_slots()
    for s in slots:
        s.enabled = True
    agg = demo_aggregator()
    secret = "PROJECT_OMEGA_KEY"

    async def mk():
        async with httpx.AsyncClient() as c:
            return await make_rule(c, agg, f"task {secret}", f"bad {secret}", "thiếu kiểm chứng")

    rule, _ = asyncio.run(mk())
    assert rule is not None and secret not in rule.guidance
    p = tmp_path / "mem.json"
    MemoryStore(str(p)).add_rule(rule)
    assert secret not in p.read_text(encoding="utf-8")        # KHÔNG bao giờ persist raw

    _, trail = asyncio.run(theosis("giải thích vòng lặp", slots, agg, max_rounds=1,
                                   memory_rules=["Luôn nêu điều kiện dừng"]))
    assert trail["memory"] == ["Luôn nêu điều kiện dừng"]
    assert any("[ký-ức]" in a for a in trail["fanout"].values())  # ký ức tới được fan-out


def test_learn_and_use_memory_via_api(tmp_path, monkeypatch):
    import json as _json

    monkeypatch.setenv("THEOSIS_DEMO", "1")
    monkeypatch.setenv("THEOSIS_MEMORY", str(tmp_path / "mem.json"))
    monkeypatch.setenv("THEOSIS_METRICS", str(tmp_path / "met.jsonl"))
    from fastapi.testclient import TestClient

    from theosis import server
    from theosis.memory import Rule

    c = TestClient(server.app)
    server.STORE.clear()

    secret = "INTERNAL_CODE_777"
    r = c.post("/api/learn", json={"request": f"x {secret}", "bad_answer": f"y {secret}", "correction": "z"})
    assert r.status_code == 200 and secret not in r.text       # /api/learn content-free
    assert len(server.STORE.all_rules()) == 1

    server.STORE.clear()
    server.STORE.add_rule(Rule(guidance="Nêu điều kiện dừng", task_type="code", keywords=["đệ quy"]))
    lines = [_json.loads(x) for x in c.post(
        "/api/run", json={"prompt": "giải thích đệ quy", "enabled_slots": ["opus", "gpt"], "use_memory": True}
    ).text.splitlines() if x.strip()]
    assert "memory" in [e["type"] for e in lines]              # event memory phát ra
    assert server.STORE.all_rules()[0].uses == 1              # bump_uses chạy


# ── V3 / Phase C: auto-learn từ executor-fail + eval harness ─────────────────
def test_pick_failure_selects_failed_answer():
    from theosis.core import _pick_failure

    ev = {"a": {"status": "pass"}, "b": {"status": "fail", "summary": "2+2=5 sai"}, "c": {"status": "na"}}
    ans = {"a": "A", "b": "B", "c": "C"}
    assert _pick_failure(ans, ev) == ("b", "B", "2+2=5 sai")
    assert _pick_failure(ans, {"a": {"status": "pass"}}) is None
    assert _pick_failure(ans, {}) is None


def test_auto_learn_from_executor_fail_is_content_free():
    slots = demo_slots()
    for s in slots:
        s.enabled = True
    agg = demo_aggregator()
    secret = "OMEGA_SECRET_KEY"
    # max_rounds=0 -> fan-out = câu cuối; số học sai (mock echo prompt) sống tới cuối
    _, trail = asyncio.run(theosis(f"Có phải 2 + 2 = 5 không? {secret}", slots, agg,
                                   max_rounds=0, use_executor=True, auto_learn=True))
    assert trail["learned"] and len(trail["learned"]) == 1
    rule = trail["learned"][0]
    assert rule["source"] == "auto:executor_fail"
    assert secret not in rule["guidance"] and all(secret not in k for k in rule["keywords"])

    # câu đúng -> không học; auto_learn tắt -> không học
    _, ok = asyncio.run(theosis("Có phải 2 + 2 = 4 không?", slots, agg, max_rounds=0, use_executor=True, auto_learn=True))
    assert ok["learned"] is None
    _, off = asyncio.run(theosis("Có phải 2 + 2 = 5 không?", slots, agg, max_rounds=0, use_executor=True))
    assert off["learned"] is None


def test_eval_harness_runs_both_arms():
    from theosis.eval import EvalTask, _run_metrics, run_eval

    slots = demo_slots()
    for s in slots:
        s.enabled = True
    agg = demo_aggregator()
    # merge mock chứa "hợp nhất" -> check kiểm chứng được cả hai chiều
    tasks = [EvalTask("câu A", expect_contains="hợp nhất"),
             EvalTask("câu B", expect_contains="KHONG_CO_TRONG_FINAL")]
    rep = asyncio.run(run_eval(tasks, slots, agg, store=None, rounds=1, use_executor=False, trials=2))
    assert rep.n_tasks == 2 and rep.n_trials == 2 and rep.has_ground_truth
    assert rep.baseline["n"] == 4 and rep.treatment["n"] == 4
    assert rep.baseline["check_pass_rate"] == 0.5     # A pass, B fail
    assert "check_pass_rate" in rep.deltas and isinstance(rep.verdict, str)

    m = _run_metrics(EvalTask("x", expect_contains="abc"), "... abc ...",
                     {"scores": {"a": 0.8}, "cost": {"total_tokens": 5}, "rounds": [1]})
    assert m["check"] is True and m["score"] == 0.8
    assert _run_metrics(EvalTask("x"), "y", {})["check"] is None


def test_eval_load_tasks(tmp_path):
    from theosis.eval import load_tasks

    p = tmp_path / "t.json"
    p.write_text('{"tasks":[{"prompt":"a","expect_contains":"x"},{"prompt":"b","task_type":"code"}]}', encoding="utf-8")
    tasks = load_tasks(str(p))
    assert len(tasks) == 2 and tasks[0].expect_contains == "x" and tasks[1].task_type == "code"


def test_auto_learn_persists_via_api(tmp_path, monkeypatch):
    import json as _json

    monkeypatch.setenv("THEOSIS_DEMO", "1")
    monkeypatch.setenv("THEOSIS_MEMORY", str(tmp_path / "mem.json"))
    monkeypatch.setenv("THEOSIS_METRICS", str(tmp_path / "met.jsonl"))
    from fastapi.testclient import TestClient

    from theosis import server

    c = TestClient(server.app)
    server.STORE.clear()
    lines = [_json.loads(x) for x in c.post("/api/run", json={
        "prompt": "Có phải 2 + 2 = 5 không?", "enabled_slots": ["opus", "gpt"],
        "use_executor": True, "max_rounds": 0, "auto_learn": True,
    }).text.splitlines() if x.strip()]
    assert "learned" in [e["type"] for e in lines]
    assert len(server.STORE.all_rules()) == 1
    assert server.STORE.all_rules()[0].source == "auto:executor_fail"


# ── V3 / Phase D: review/verify rule ────────────────────────────────────────
def test_memory_verify_edit_and_verified_only(tmp_path):
    from theosis.memory import MemoryStore, Rule

    s = MemoryStore(str(tmp_path / "m.json"))
    a = s.add_rule(Rule(guidance="A", task_type="code", keywords=["x"], verified=True))
    b = s.add_rule(Rule(guidance="B", task_type="code", keywords=["x"], verified=False))

    assert len(s.relevant(task_type="code")) == 2
    assert [r.id for r in s.relevant(task_type="code", verified_only=True)] == [a.id]
    assert s.set_verified(b.id, True) and len(s.relevant(task_type="code", verified_only=True)) == 2
    assert s.set_verified("nope") is False
    assert s.update_guidance(a.id, "A sửa") and any(r.guidance == "A sửa" for r in s.all_rules())
    assert s.update_guidance(a.id, "   ") is False
    # verified persist qua reload
    assert all(r.verified for r in MemoryStore(str(tmp_path / "m.json")).all_rules())


def test_verify_gate_and_endpoints_via_api(tmp_path, monkeypatch):
    import json as _json

    monkeypatch.setenv("THEOSIS_DEMO", "1")
    monkeypatch.setenv("THEOSIS_MEMORY", str(tmp_path / "mem.json"))
    monkeypatch.setenv("THEOSIS_METRICS", str(tmp_path / "met.jsonl"))
    from fastapi.testclient import TestClient

    from theosis import server
    from theosis.memory import Rule

    c = TestClient(server.app)
    S = server.STORE
    S.clear()

    # teach tay -> verified True; auto-learn -> verified False
    assert c.post("/api/learn", json={"request": "x", "bad_answer": "y", "correction": "z"}).json()["verified"] is True
    S.clear()
    c.post("/api/run", json={"prompt": "Có phải 2 + 2 = 5 không?", "enabled_slots": ["opus", "gpt"],
                             "use_executor": True, "max_rounds": 0, "auto_learn": True})
    rid = S.all_rules()[0].id
    assert S.all_rules()[0].verified is False

    # endpoints verify + edit
    assert c.post(f"/api/memory/{rid}/verify", json={"verified": True}).status_code == 200
    assert S.all_rules()[0].verified is True
    assert c.post("/api/memory/nope/verify", json={"verified": True}).status_code == 404
    assert c.patch(f"/api/memory/{rid}", json={"guidance": "đã sửa"}).status_code == 200
    assert c.patch(f"/api/memory/{rid}", json={"guidance": "  "}).status_code == 400

    # GATE: rule chưa duyệt bị chặn khi verified_only
    S.clear()
    S.add_rule(Rule(guidance="chưa duyệt", task_type="code", keywords=["đệ quy"], verified=False))

    def run(vo):
        body = {"prompt": "giải thích đệ quy", "enabled_slots": ["opus", "gpt"],
                "use_memory": True, "verified_only": vo}
        return [_json.loads(x) for x in c.post("/api/run", json=body).text.splitlines() if x.strip()]

    assert "memory" in [e["type"] for e in run(False)]      # off: tiêm
    assert "memory" not in [e["type"] for e in run(True)]   # on: chặn
    S.set_verified(S.all_rules()[0].id, True)
    assert "memory" in [e["type"] for e in run(True)]        # duyệt xong: tiêm lại


def test_config_exposes_all_settings(monkeypatch):
    monkeypatch.setenv("THEOSIS_DEMO", "1")
    from theosis.config import DEFAULT_SETTINGS, load_config

    _, _, settings = load_config()
    for k in ("use_router", "use_memory", "auto_learn", "verified_only", "strategy"):
        assert k in settings
    assert set(DEFAULT_SETTINGS).issubset(set(settings))


# ── V3 / Phase E: auto-learn đa nguồn ───────────────────────────────────────
def test_learn_signal_priority_and_branches():
    from theosis.core import _learn_signal

    ans = {"a": "A", "b": "B"}
    # 1) executor_fail thắng tất cả
    s = _learn_signal(ans, {"b": {"status": "fail", "summary": "x"}}, {"a": 0.1, "b": 0.1}, {}, {"rounds": [1, 2]})
    assert s[0] == "auto:executor_fail" and s[1] == "b"
    # 2) low_confidence khi điểm tốt nhất < ngưỡng
    s = _learn_signal(ans, {}, {"a": 0.2, "b": 0.1}, {"a": [{"critique": "tệ"}]}, {})
    assert s[0] == "auto:low_confidence" and s[1] == "a"
    # 3) no_converge khi chạy hết ≥2 vòng mà chưa hội tụ
    s = _learn_signal(ans, {}, {"a": 0.5}, {}, {"rounds": [1, 2], "stopped_reason": None})
    assert s[0] == "auto:no_converge"
    # không tín hiệu
    assert _learn_signal(ans, {}, {"a": 0.5}, {}, {"rounds": [1, 2], "stopped_reason": "converged"}) is None
    assert _learn_signal(ans, {}, {"a": 0.5}, {}, {"rounds": [1], "stopped_reason": None}) is None
    assert _learn_signal(ans, {}, {}, {}, {}) is None


def test_auto_learn_three_sources_content_free():
    slots = demo_slots()
    for s in slots:
        s.enabled = True
    agg = demo_aggregator()
    sec = "SECRET_MULTI_42"

    _, t1 = asyncio.run(theosis(f"Có phải 2 + 2 = 5? {sec}", slots, agg, max_rounds=0, use_executor=True, auto_learn=True))
    _, t2 = asyncio.run(theosis(f"giải thích đệ quy {sec}", slots, agg, max_rounds=0, auto_learn=True, low_confidence_threshold=0.6))
    _, t3 = asyncio.run(theosis(f"đạo đức AI {sec}", slots, agg, max_rounds=2, converge_threshold=2.0, auto_learn=True))
    assert t1["learned"][0]["source"] == "auto:executor_fail"
    assert t2["learned"][0]["source"] == "auto:low_confidence"
    assert t3["learned"][0]["source"] == "auto:no_converge"
    assert all(sec not in t["learned"][0]["guidance"] for t in (t1, t2, t3))   # content-free


def test_memory_add_rule_dedup_bumps_uses(tmp_path):
    from theosis.memory import MemoryStore, Rule

    s = MemoryStore(str(tmp_path / "m.json"))
    s.add_rule(Rule(guidance="X", task_type="code"), dedup=True)
    s.add_rule(Rule(guidance="X", task_type="code"), dedup=True)   # trùng -> bump
    assert len(s.all_rules()) == 1 and s.all_rules()[0].uses == 1
    s.add_rule(Rule(guidance="Y"), dedup=True)
    assert len(s.all_rules()) == 2
    s.add_rule(Rule(guidance="X"))                                  # không dedup -> thêm
    assert len(s.all_rules()) == 3


# ── V3 / Phase F: auto-demote rule (học ngược từ eval) ──────────────────────
def test_memory_demote_excludes_restores_and_backcompat(tmp_path):
    import json as _json

    from theosis.memory import MemoryStore, Rule

    s = MemoryStore(str(tmp_path / "m.json"))
    a = s.add_rule(Rule(guidance="A", task_type="code", keywords=["x"], verified=True))
    b = s.add_rule(Rule(guidance="B", task_type="code", keywords=["x"], verified=True))
    assert a.score == 0.0 and a.demoted is False
    assert s.set_demoted(b.id, True, score=-0.2)
    assert [r.id for r in s.relevant(task_type="code")] == [a.id]      # demoted bị loại
    assert len(s.all_rules()) == 2                                     # vẫn giữ để review
    assert b.id not in [r.id for r in s.relevant(task_type="code", verified_only=False)]
    assert s.set_demoted(b.id, False) and len(s.relevant(task_type="code")) == 2
    assert s.set_score(a.id, 0.3) and next(r for r in s.all_rules() if r.id == a.id).score == 0.3
    assert s.set_demoted("nope") is False

    # backward-compat: file cũ không có score/demoted
    old = tmp_path / "old.json"
    old.write_text(_json.dumps({"rules": [{"guidance": "Cũ", "task_type": "code", "keywords": [],
                  "source": "manual", "id": "r_old", "created_at": "2025-01-01T00:00:00Z",
                  "uses": 2, "verified": True}]}), encoding="utf-8")
    r = MemoryStore(str(old)).all_rules()[0]
    assert r.guidance == "Cũ" and r.score == 0.0 and r.demoted is False and r.uses == 2


def test_eval_rules_machinery_and_auto_demote():
    from theosis.eval import EvalTask, RuleEval, RuleEvalReport, auto_demote, eval_rules
    from theosis.memory import MemoryStore, Rule
    import tempfile

    slots = demo_slots()
    for s in slots:
        s.enabled = True
    agg = demo_aggregator()
    d = tempfile.mkdtemp()
    st = MemoryStore(d + "/m.json")
    st.add_rule(Rule(guidance="r1", task_type="reasoning", keywords=["a"]))
    st.add_rule(Rule(guidance="r2", task_type="reasoning", keywords=["b"]))
    rep = asyncio.run(eval_rules([EvalTask("câu X", expect_contains="hợp nhất")], slots, agg, st,
                                 rounds=1, use_executor=False, trials=1))
    assert rep.n_rules == 2 and len(rep.per_rule) == 2
    assert all(abs(r.marginal or 0) < 0.01 for r in rep.per_rule)     # mock: ~0, không demote bừa
    assert auto_demote(st, rep, threshold=-0.05) == []

    # auto_demote hạ đúng rule có hại (report tổng hợp)
    st2 = MemoryStore(d + "/m2.json")
    A = st2.add_rule(Rule(guidance="tốt", task_type="code", keywords=["k"]))
    B = st2.add_rule(Rule(guidance="hại", task_type="code", keywords=["k"]))
    synth = RuleEvalReport(n_rules=2, n_tasks=3, n_trials=2, has_ground_truth=True, baseline={}, per_rule=[
        RuleEval(id=A.id, guidance="tốt", source="manual", delta_check=0.1, delta_score=None, marginal=0.1, demoted=False, n=6),
        RuleEval(id=B.id, guidance="hại", source="auto:low_confidence", delta_check=-0.2, delta_score=None, marginal=-0.2, demoted=False, n=6),
    ])
    assert auto_demote(st2, synth, threshold=-0.05) == [B.id]
    rb = next(r for r in st2.all_rules() if r.id == B.id)
    ra = next(r for r in st2.all_rules() if r.id == A.id)
    assert rb.demoted is True and rb.score == -0.2
    assert ra.demoted is False and ra.score == 0.1            # cập nhật score, không demote
    assert B.id not in [r.id for r in st2.relevant(task_type="code")]


def test_demote_endpoints_via_api(tmp_path, monkeypatch):
    monkeypatch.setenv("THEOSIS_DEMO", "1")
    monkeypatch.setenv("THEOSIS_MEMORY", str(tmp_path / "mem.json"))
    from fastapi.testclient import TestClient

    from theosis import server
    from theosis.memory import Rule

    c = TestClient(server.app)
    S = server.STORE
    S.clear()
    r = S.add_rule(Rule(guidance="X", task_type="code", keywords=["k"]))
    assert c.post(f"/api/memory/{r.id}/demote", json={"demoted": True, "score": -0.3}).status_code == 200
    assert S.all_rules()[0].demoted is True and S.all_rules()[0].score == -0.3
    assert S.relevant(task_type="code") == []
    assert c.post(f"/api/memory/{r.id}/demote", json={"demoted": False}).status_code == 200
    assert len(S.relevant(task_type="code")) == 1
    assert c.post("/api/memory/nope/demote", json={"demoted": True}).status_code == 404
    assert all(k in c.get("/api/memory").json()["rules"][0] for k in ("score", "demoted"))


# ── V3 / Phase G: thống kê rule theo thời gian ──────────────────────────────
def test_rule_stats(tmp_path):
    from theosis.memory import MemoryStore, Rule
    from theosis.stats import rule_stats

    s = MemoryStore(str(tmp_path / "m.json"))
    a = s.add_rule(Rule(guidance="A", source="manual", verified=True, created_at="2026-06-20T01:00:00Z"))
    s.bump_uses([a.id])
    s.bump_uses([a.id])
    b = s.add_rule(Rule(guidance="B", source="auto:executor_fail", created_at="2026-06-21T01:00:00Z"))
    c = s.add_rule(Rule(guidance="C", source="auto:low_confidence", created_at="2026-06-21T02:00:00Z"))
    s.set_demoted(c.id, True, score=-0.2)
    s.set_score(b.id, 0.1)
    s.bump_uses([b.id])

    rs = rule_stats(s)
    assert rs["total"] == 3
    assert rs["verified"] == 1 and rs["unverified"] == 2 and rs["demoted"] == 1
    assert rs["helpful"] == 1 and rs["harmful"] == 1 and rs["neutral"] == 1
    assert rs["by_source"] == {"manual": 1, "auto:executor_fail": 1, "auto:low_confidence": 1}
    assert rs["total_uses"] == 3
    assert rs["top_used"][0]["uses"] == 2
    assert [x["date"] for x in rs["created_by_day"]] == ["2026-06-20", "2026-06-21"]
    assert rs["created_by_day"][1]["count"] == 2


def test_run_stats(tmp_path):
    import json as _json

    from theosis.stats import run_stats

    p = tmp_path / "met.jsonl"
    recs = [
        {"ts": "2026-06-21T03:00:00Z", "task_type": "code", "cost": {"total_tokens": 100}, "scores": [0.8, 0.6]},
        {"ts": "2026-06-21T04:00:00Z", "task_type": "code", "cost": {"total_tokens": 200}, "scores": [0.4]},
        {"ts": "2026-06-22T01:00:00Z", "task_type": "factual", "cost": {"total_tokens": 50}, "scores": [0.9]},
    ]
    p.write_text("\n".join(_json.dumps(x) for x in recs) + "\n", encoding="utf-8")
    rs = run_stats(str(p))
    assert rs["total_runs"] == 3
    assert rs["total_tokens"] == 350 and rs["avg_tokens"] == 117
    assert rs["task_types"] == {"code": 2, "factual": 1}
    assert {x["date"]: x["count"] for x in rs["runs_by_day"]} == {"2026-06-21": 2, "2026-06-22": 1}
    assert abs(rs["avg_score_by_day"][0]["avg"] - 0.55) < 0.01
    # thiếu file -> rỗng an toàn
    assert run_stats(str(tmp_path / "khongco.jsonl"))["total_runs"] == 0


def test_stats_endpoint(tmp_path, monkeypatch):
    monkeypatch.setenv("THEOSIS_DEMO", "1")
    monkeypatch.setenv("THEOSIS_MEMORY", str(tmp_path / "mem.json"))
    monkeypatch.setenv("THEOSIS_METRICS", str(tmp_path / "met.jsonl"))
    from fastapi.testclient import TestClient

    from theosis import server
    from theosis.memory import Rule

    c = TestClient(server.app)
    server.STORE.clear()
    d = c.get("/api/stats").json()
    assert set(d.keys()) == {"rules", "runs"}
    assert d["rules"]["total"] == 0 and d["runs"]["total_runs"] == 0

    server.STORE.add_rule(Rule(guidance="đệ quy cần điều kiện dừng", task_type="reasoning", keywords=["đệ quy"], verified=True))
    c.post("/api/run", json={"prompt": "giải thích đệ quy", "enabled_slots": ["opus", "gpt"], "use_memory": True})
    d2 = c.get("/api/stats").json()
    assert d2["rules"]["total"] == 1 and d2["runs"]["total_runs"] >= 1


# ── Plumbing file: tri giác (ingest) + kho file + endpoints ──────────────────
def test_ingest_text_formats_and_fallbacks():
    import json as _json

    from theosis.ingest import format_attachments, ingest_bytes

    sec = "NOI_DUNG_999"
    assert ingest_bytes(f"hi {sec}".encode(), "a.txt").text.find(sec) >= 0
    assert ingest_bytes(b"print(1)", "x.py").kind == "code"
    rj = ingest_bytes(_json.dumps({"k": sec}).encode(), "d.json")
    assert rj.kind == "json" and sec in rj.text
    rc = ingest_bytes(f"a,b\n1,{sec}".encode(), "t.csv")
    assert rc.kind == "csv" and "a | b" in rc.text and sec in rc.text
    assert "a | b" in ingest_bytes(f"a\tb\n1\t{sec}".encode(), "t.tsv").text
    assert ingest_bytes(b"k: v", "c.yaml").kind == "yaml"
    ri = ingest_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 40, "p.png")
    assert ri.kind == "image" and ri.needs_vision
    big = ingest_bytes(("x" * 25000).encode(), "big.txt")
    assert big.truncated and big.chars <= 20100
    att = format_attachments([ingest_bytes(b"hi", "a.txt"), ingest_bytes(b"x,y\n1,2", "b.csv")])
    assert "[TÀI LIỆU ĐÍNH KÈM]" in att and "a.txt" in att and "[HẾT TÀI LIỆU]" in att


def test_ingest_documents():
    import io

    pytest.importorskip("docx")
    pytest.importorskip("openpyxl")
    pytest.importorskip("pypdf")
    import docx
    import openpyxl
    from pypdf import PdfWriter

    from theosis.ingest import ingest_bytes

    sec = "DOC_SECRET_321"
    d = docx.Document()
    d.add_paragraph(f"Đoạn {sec}")
    b = io.BytesIO()
    d.save(b)
    assert sec in ingest_bytes(b.getvalue(), "f.docx").text

    wb = openpyxl.Workbook()
    wb.active["A1"] = "h"
    wb.active["B1"] = sec
    b2 = io.BytesIO()
    wb.save(b2)
    rx = ingest_bytes(b2.getvalue(), "s.xlsx")
    assert "Sheet:" in rx.text and sec in rx.text

    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    b3 = io.BytesIO()
    w.write(b3)
    assert ingest_bytes(b3.getvalue(), "f.pdf").meta.get("pages") == 1


def test_extract_code_blocks():
    from theosis.ingest import extract_code_blocks

    blocks = extract_code_blocks("x:\n```python\nprint('hi')\n```\ny:\n```sql\nSELECT 1;\n```")
    assert len(blocks) == 2 and blocks[0]["ext"] == ".py" and blocks[1]["ext"] == ".sql"
    assert "print('hi')" in blocks[0]["content"]
    assert extract_code_blocks("không có code") == []


def test_filestore(tmp_path):
    from theosis.files import FileStore

    sec = "FS_SECRET_654"
    fs = FileStore(base=str(tmp_path))
    m1 = fs.add_upload("note.txt", f"x {sec}".encode())
    m2 = fs.add_upload("data.csv", b"a,b\n1,2")
    assert m1["token"].startswith("u_") and m1["kind"] == "text" and sec in m1["preview"]
    ctx = fs.text_for([m1["token"], m2["token"]])
    assert "note.txt" in ctx and "data.csv" in ctx and sec in ctx
    assert fs.text_for([]) == "" and "note.txt" in fs.text_for([m1["token"], "u_x"])
    o = fs.add_output("s.py", b"print(1)")
    assert o["url"] == f"/api/files/{o['token']}"
    assert fs.read(o["token"])[0] == b"print(1)"
    assert fs.read("x_none") is None
    with pytest.raises(ValueError):
        fs.add_upload("big.bin", b"x" * (26 * 1024 * 1024))
    fs.clear()
    assert fs.text_for([m1["token"]]) == "" and fs.read(o["token"]) is None


def test_upload_run_download_endpoints(tmp_path, monkeypatch):
    import json as _json

    monkeypatch.setenv("THEOSIS_DEMO", "1")
    monkeypatch.setenv("THEOSIS_MEMORY", str(tmp_path / "mem.json"))
    monkeypatch.setenv("THEOSIS_FILES", str(tmp_path / "files"))
    from fastapi.testclient import TestClient

    from theosis import server

    c = TestClient(server.app)
    server.FILES.clear()
    sec = "ENDPOINT_SECRET_111"

    up = c.post("/api/upload", files=[
        ("files", ("note.txt", f"bí mật {sec}".encode(), "text/plain")),
        ("files", ("data.csv", b"ten,tuoi\nAn,30", "text/csv")),
    ])
    assert up.status_code == 200 and len(up.json()["files"]) == 2
    toks = [f["token"] for f in up.json()["files"]]

    ev = [_json.loads(x) for x in c.post("/api/run", json={
        "prompt": "Tóm tắt", "enabled_slots": ["opus", "gpt"], "file_tokens": toks,
    }).text.splitlines() if x.strip()]
    assert any(e["type"] == "attachments" and len(e["files"]) == 2 for e in ev)
    # council THẤY nội dung file (lọt vào fan-out — chia chung)
    assert any(sec in _json.dumps(e, ensure_ascii=False) for e in ev if e["type"] == "fanout_done")

    dl = c.get(f"/api/files/{toks[0]}")
    assert dl.status_code == 200 and sec.encode() in dl.content
    assert c.get("/api/files/x_none").status_code == 404


# ── Council đa-lượt (multi-turn) ─────────────────────────────────────────────
def test_prepare_history_window_and_cap():
    from theosis.core import _history_messages, _prepare_history

    h = [{"role": "user", "content": "a"}, {"role": "assistant", "content": "b"},
         {"role": "bogus", "content": "x"}, {"role": "user", "content": ""},
         {"role": "assistant", "content": "c"}]
    assert [m["role"] for m in _prepare_history(h)] == ["user", "assistant", "assistant"]
    assert len(_prepare_history([{"role": "user", "content": str(i)} for i in range(20)], max_turns=5)) == 5
    assert _prepare_history([{"role": "user", "content": "x" * 5000}], max_chars=100)[0]["content"].endswith("…[cắt]")
    assert _prepare_history(None) == []
    assert _history_messages([{"role": "user", "content": "ok"}]) == [{"role": "user", "content": "ok"}]


def test_multi_turn_council_sees_history():
    slots = demo_slots()
    for s in slots:
        s.enabled = True
    agg = demo_aggregator()
    hist = [{"role": "user", "content": "Tên tôi là Minh"}, {"role": "assistant", "content": "Chào Minh"}]
    _, t = asyncio.run(theosis("Tôi tên gì?", slots, agg, max_rounds=0, history=hist))
    assert all("[đa-lượt]" in v for v in t["fanout"].values())   # marker => history tới council
    _, t0 = asyncio.run(theosis("Xin chào", slots, agg, max_rounds=0))
    assert not any("[đa-lượt]" in v for v in t0["fanout"].values())
    # đi hết vòng audit + merge với history không vỡ
    fm, _ = asyncio.run(theosis("Câu hỏi", slots, agg, max_rounds=1, history=hist))
    assert isinstance(fm, str) and fm


def test_ask_slot_orders_history_before_current_user():
    import httpx

    import theosis.core as core

    slots = demo_slots()
    for s in slots:
        s.enabled = True

    captured = {}
    orig = core._call

    async def spy(client, slot, messages, **kw):
        captured["m"] = messages
        return await orig(client, slot, messages, **kw)

    async def go():
        core._call = spy
        try:
            async with httpx.AsyncClient() as c:
                await core.ask_slot(c, slots[0], "MỚI",
                                    history=[{"role": "user", "content": "cũ"}, {"role": "assistant", "content": "đáp"}])
        finally:
            core._call = orig

    asyncio.run(go())
    roles = [m["role"] for m in captured["m"]]
    assert roles[-3:] == ["user", "assistant", "user"] and captured["m"][-1]["content"] == "MỚI"


def test_split_history_helper():
    from theosis.server import Msg, _split_history

    msgs = [Msg(role="system", content="s"), Msg(role="user", content="t1"),
            Msg(role="assistant", content="a1"), Msg(role="user", content="t2")]
    h, u = _split_history(msgs)
    assert u == "t2" and len(h) == 3 and h[-1]["content"] == "a1"
    assert _split_history([Msg(role="assistant", content="x")]) == ([], "")


def test_endpoints_thread_history(tmp_path, monkeypatch):
    import json as _json

    monkeypatch.setenv("THEOSIS_DEMO", "1")
    monkeypatch.setenv("THEOSIS_MEMORY", str(tmp_path / "mem.json"))
    from fastapi.testclient import TestClient

    from theosis import server

    c = TestClient(server.app)

    # /api/run với history → council thấy ([đa-lượt])
    ev = [_json.loads(x) for x in c.post("/api/run", json={
        "prompt": "Tôi tên gì?", "enabled_slots": ["opus", "gpt"],
        "history": [{"role": "user", "content": "Tên tôi là Minh"}, {"role": "assistant", "content": "Chào Minh"}],
    }).text.splitlines() if x.strip()]
    assert "[đa-lượt]" in _json.dumps([e for e in ev if e["type"] == "fanout_done"], ensure_ascii=False)

    # /v1 đa lượt (non-stream + stream)
    r = c.post("/v1/chat/completions", json={"model": "theosis-v1", "messages": [
        {"role": "user", "content": "Nhớ 42"}, {"role": "assistant", "content": "OK"}, {"role": "user", "content": "Số?"}]})
    assert r.status_code == 200 and "choices" in r.json()
    rs = c.post("/v1/chat/completions", json={"model": "theosis-v1", "stream": True, "messages": [
        {"role": "user", "content": "A"}, {"role": "assistant", "content": "B"}, {"role": "user", "content": "C"}]})
    assert rs.status_code == 200 and "[DONE]" in rs.text


# ── Tóm tắt hội thoại dài (summarize_history) ────────────────────────────────
def test_prepare_history_smart_summarize():
    import httpx

    from theosis.core import _prepare_history_smart

    slots = demo_slots()
    for s in slots:
        s.enabled = True
    agg = demo_aggregator()
    H = [{"role": "user" if i % 2 == 0 else "assistant", "content": f"lượt {i} XYZ{i}"} for i in range(12)]

    async def go():
        async with httpx.AsyncClient() as c:
            win, res0 = await _prepare_history_smart(c, agg, H, 4, False)
            assert len(win) == 4 and res0 is None and win[0]["content"] == "lượt 8 XYZ8"
            eff, res1 = await _prepare_history_smart(c, agg, H, 4, True)
            assert eff[0]["role"] == "system" and "TÓM TẮT" in eff[0]["content"]
            assert "(mock·tóm tắt)" in eff[0]["content"]
            assert len(eff) == 5 and eff[-1]["content"] == "lượt 11 XYZ11" and res1 is not None
            eff2, res2 = await _prepare_history_smart(c, agg, H[:3], 4, True)
            assert res2 is None and len(eff2) == 3   # ngắn -> không tốn call summary

    asyncio.run(go())


def test_summarize_end_to_end_and_server_flag(tmp_path, monkeypatch):
    import json as _json

    slots = demo_slots()
    for s in slots:
        s.enabled = True
    agg = demo_aggregator()
    H = [{"role": "user" if i % 2 == 0 else "assistant", "content": f"lượt {i}"} for i in range(12)]
    _, t = asyncio.run(theosis("Câu mới", slots, agg, max_rounds=0, history=H,
                               max_history_turns=4, summarize_history=True))
    assert all("[đa-lượt]" in v for v in t["fanout"].values())

    monkeypatch.setenv("THEOSIS_DEMO", "1")
    monkeypatch.setenv("THEOSIS_MEMORY", str(tmp_path / "mem.json"))
    from fastapi.testclient import TestClient

    from theosis import server

    c = TestClient(server.app)
    assert "summarize_history" in c.get("/api/config").json()["settings"]
    ev = [_json.loads(x) for x in c.post("/api/run", json={
        "prompt": "Câu mới", "enabled_slots": ["opus", "gpt"], "history": H, "summarize_history": True,
    }).text.splitlines() if x.strip()]
    assert any(e["type"] == "done" for e in ev)
